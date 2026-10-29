"""把 Markdown 报告转成 Word：标题、表格、列表、图片、加粗。"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt
from docx.oxml.ns import qn

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def _set_cn_font(document: Document):
    style = document.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _add_runs(paragraph, text: str):
    # 链接只保留文字；行内代码去掉反引号
    text = LINK_RE.sub(lambda m: m.group(1), text)
    text = text.replace("`", "")

    pos = 0
    for match in BOLD_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        run = paragraph.add_run(match.group(1))
        run.bold = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _flush_table(document: Document, rows: list):
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = ""
            value = row[j] if j < len(row) else ""
            _add_runs(cell.paragraphs[0], value)
    document.add_paragraph()


def markdown_to_docx(md_text: str, output_path: str, base_dir: str = ".") -> str:
    document = Document()
    _set_cn_font(document)

    table_rows = []
    in_code = False

    def flush():
        nonlocal table_rows
        _flush_table(document, table_rows)
        table_rows = []

    for raw in md_text.splitlines():
        line = raw.rstrip()

        if line.startswith("```"):
            flush()
            in_code = not in_code
            continue

        if in_code:
            para = document.add_paragraph(line)
            for run in para.runs:
                run.font.name = "Consolas"
            continue

        if not line.strip():
            flush()
            continue

        if line.startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if cells and all(set(c) <= set("-: ") for c in cells if c):
                continue
            table_rows.append(cells)
            continue

        flush()

        image = IMAGE_RE.match(line.strip())
        if image:
            path = Path(base_dir) / image.group(2)
            if path.exists():
                try:
                    document.add_picture(str(path), width=Cm(15))
                    document.add_paragraph()
                except Exception:
                    pass
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading:
            document.add_heading(heading.group(2).strip(), level=min(len(heading.group(1)), 4))
            continue

        if line.strip() in ("---", "***"):
            continue

        bullet = re.match(r"^\s*[-*]\s+(.*)$", line)
        if bullet:
            _add_runs(document.add_paragraph(style="List Bullet"), bullet.group(1))
            continue

        _add_runs(document.add_paragraph(), line.strip())

    flush()
    document.save(output_path)
    return output_path
